import json, pathlib

# Check raw/detail/111 for Part A fields
d111 = json.load(open('raw/detail/111/getWebsiteStandardDetails.json', encoding='utf-8'))['data']
print('=== raw/detail/111 Part A fields ===')
for f in ['ministryName','groupName','subGroupName','subSubGroupName','certificationName']:
    print(f'  {f}: {repr(d111.get(f))}')

print()

# Check mandatory standards with blank QCO / Ministry
MAND_BLANK = ['4003','7823','8040','8195','10274','10335','10951','12837',
              '12840','18590','25777','25982','27507','27815','27935','27936',
              '28328','29096','30417','30554','34001','34006']

print('=== Mandatory with blank QCO - raw ministryName ===')
for pk in MAND_BLANK:
    fp = pathlib.Path(f'raw/detail/{pk}/getWebsiteStandardDetails.json')
    if fp.exists():
        d = json.load(open(fp, encoding='utf-8'))['data']
        mn = d.get('ministryName') or ''
        cn = d.get('certificationName') or ''
        sn = d.get('standardNumber','')
        print(f'  pk={pk} {sn} cert={repr(cn)} ministry={repr(mn)}')
    else:
        print(f'  pk={pk} NOT IN CRAWLED 400')

print()

# Check Standards Master - does isStatus_label column exist?
import openpyxl
wb = openpyxl.load_workbook('sih_bis_mvp_400_hardened.xlsx')
ws = wb['Standards Master (400)']
headers = [c.value for c in ws[1]]
print('isStatus_label in headers?', 'isStatus_label' in headers)
print('withdrawStatus in headers?', 'withdrawStatus' in headers)
print('isStatus in headers?', 'isStatus' in headers)
print('Status col value sample (rows 2-5):')
if 'Status' in headers:
    si = headers.index('Status') + 1
    for row in ws.iter_rows(min_row=2, max_row=5, values_only=True):
        print(f'  Status={row[si-1]}')
