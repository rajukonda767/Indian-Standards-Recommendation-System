import openpyxl, json, pathlib
from collections import Counter

wb = openpyxl.load_workbook('sih_bis_mvp_400_hardened_v2.xlsx')
ws = wb['Standards Master (400)']
headers = [c.value for c in ws[1]]
print('Cols:', len(headers))
print('New cols (21 onwards):', headers[20:])

# Check if sn_to_pk mapping works
RAW = pathlib.Path('raw/detail')
sn_to_pk = {}
for pk_dir in RAW.iterdir():
    fp = pk_dir / 'getWebsiteStandardDetails.json'
    if fp.exists():
        d = json.load(open(fp, encoding='utf-8')).get('data', {})
        sn = (d.get('standardNumber') or '').strip()
        if sn:
            sn_to_pk[sn] = pk_dir.name
print('sn_to_pk entries:', len(sn_to_pk))

# Cross-check Standard Number in master vs sn_to_pk
sn_idx = headers.index('Standard Number')
il_idx = headers.index('isStatus_label') if 'isStatus_label' in headers else None
print('First 5 standard numbers in master vs sn_to_pk:')
for row in ws.iter_rows(min_row=2, max_row=6, values_only=True):
    sn = row[sn_idx]
    pk = sn_to_pk.get(sn or '', 'MISSING')
    lbl = row[il_idx] if il_idx is not None else '?'
    print(f'  sn={sn}  pk={pk}  isStatus_label={lbl}')

# Count QCO backfill result
cert_idx = headers.index('Certification')
qco_idx = headers.index('QCO / Ministry')
mand_blank = 0
mand_filled = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    cert = str(row[cert_idx] or '').strip()
    qco  = str(row[qco_idx] or '').strip()
    if cert == 'Mandatory':
        if qco:
            mand_filled += 1
        else:
            mand_blank += 1
            if mand_blank <= 5:
                sn = row[sn_idx]
                print(f'  Mandatory+blankQCO: SN={sn}')

print(f'Mandatory with QCO filled: {mand_filled}')
print(f'Mandatory with QCO blank: {mand_blank}')

# B1: isStatus_label categories
if il_idx is not None:
    vc = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        vc[row[il_idx]] += 1
    print('isStatus_label distribution:', dict(vc))

# Ministry col check for pk=111 (IS 269:2015)
min_idx = headers.index('Ministry') if 'Ministry' in headers else None
for row in ws.iter_rows(min_row=2, values_only=True):
    sn = row[sn_idx]
    if sn and '269:2015' in str(sn):
        print(f'IS 269:2015 Ministry={row[min_idx] if min_idx else "?"}  isStatus_label={row[il_idx] if il_idx else "?"}')
        break
