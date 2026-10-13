import openpyxl

wb = openpyxl.load_workbook('sih_bis_mvp_400_hardened_v2.xlsx')
print('Sheets:', wb.sheetnames)
print()

for sn in wb.sheetnames:
    ws = wb[sn]
    headers = [c.value for c in ws[1]]
    print(f'[{sn}]  rows={ws.max_row-1}  cols={ws.max_column}')
    print(f'  headers: {headers}')
    # Show sample row 2
    row2 = [c.value for c in list(ws.iter_rows(min_row=2, max_row=2))[0]]
    print(f'  row2:    {row2[:8]}')
    print()
