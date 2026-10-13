import openpyxl, re, csv, pathlib
from collections import Counter

wb = openpyxl.load_workbook('sih_bis_mvp_400_hardened.xlsx')
ws = wb['Allied & References (Hardened)']
headers = [c.value for c in ws[1]]
print('Edge headers:', headers)

status_cnt = Counter()
junk_rows = []
not_in_index = []

for row in ws.iter_rows(min_row=2, values_only=True):
    h = {headers[i]: row[i] for i in range(len(headers))}
    rs = h.get('referenced_status', '') or ''
    ref_num = str(h.get('ref_standardNumber', '') or '').strip()
    status_cnt[rs] += 1
    if not ref_num or re.fullmatch(r'[\s0]+', ref_num):
        junk_rows.append(ref_num)
    if rs == 'NOT_IN_INDEX':
        not_in_index.append(ref_num)

print('referenced_status counts:', dict(status_cnt))
print('Junk row count:', len(junk_rows), 'samples:', junk_rows[:10])
print('NOT_IN_INDEX count:', len(not_in_index))
print('NOT_IN_INDEX sample:', not_in_index[:10])

# Load index base numbers
base_set = {}
with open('all_standards_index.csv', newline='', encoding='utf-8') as f:
    for r in csv.DictReader(f):
        b = r['base_number'].strip()
        yr = r['year'].strip()
        sn = r['standardNumber'].strip()
        ist = r['isStatus']
        if b not in base_set or (yr and int(yr or 0) > int(base_set[b]['year'] or 0)):
            base_set[b] = {'year': yr, 'isStatus': ist, 'standardNumber': sn}

print('\nIndex base_number count:', len(base_set))

def strip_to_base(sn):
    b = re.sub(r'\s*:.*$', '', sn).strip()
    b = b.replace(' ', '').replace('(', '').replace(')', '').replace('/', '')
    return b

# Check how many NOT_IN_INDEX resolve by base-only match
resolves = 0
is_looking = 0
for ref in not_in_index:
    if ref.startswith('IS') or ref.startswith('SP') or ref.startswith('IS/'):
        is_looking += 1
        base = strip_to_base(ref)
        if base in base_set:
            resolves += 1

print(f'NOT_IN_INDEX IS-looking refs: {is_looking}')
print(f'Resolve by base-only match: {resolves}')

# Show the licence varieties sheet 7 cap rows
ws_lic = wb['Licence Varieties']
lic_headers = [c.value for c in ws_lic[1]]
print('\nLicence Varieties headers:', lic_headers)
long_rows = []
pv_idx = lic_headers.index('product_varieties')
pk_idx = lic_headers.index('pk_is_id')
sn_idx = lic_headers.index('standardNumber')
for row in ws_lic.iter_rows(min_row=2, values_only=True):
    pv = str(row[pv_idx] or '')
    if len(pv) > 1000:
        long_rows.append((row[pk_idx], row[sn_idx], len(pv)))

long_rows.sort(key=lambda x: -x[2])
print(f'Rows >1000 chars in product_varieties: {len(long_rows)}')
for pk, sn, ln in long_rows[:10]:
    print(f'  pk={pk} sn={sn} len={ln}')
