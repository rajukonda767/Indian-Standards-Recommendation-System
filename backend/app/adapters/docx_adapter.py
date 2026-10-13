# -*- coding: utf-8 -*-
"""
backend/app/adapters/docx_adapter.py
====================================
Tender document adapter for Microsoft Word (.docx) files.

Reuses the core PageBlock and tier model from BaseAdapter:
- Extracts paragraphs and structured tables with hierarchical section awareness.
- Classifies boilerplate / GTC as CONTEXT_ONLY and specifications/BoQ as AUTHORITATIVE.
- Enforces strict page/element limits without silent truncation.
"""
from __future__ import annotations

import os
import re
from typing import List, Tuple, Optional
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.adapters.base import BaseAdapter, PageBlock


class DocxAdapter(BaseAdapter):
    """Adapter for Microsoft Word (.docx) tender and specification files."""

    MAX_ELEMENTS = 10000  # Hard ceiling on total paragraphs + table rows

    def can_handle(self, file_path: str) -> bool:
        ext = os.path.splitext(file_path)[1].lower()
        return ext in (".docx",)

    def extract_blocks(self, file_path: str) -> List[PageBlock]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        doc = docx.Document(file_path)
        blocks: List[PageBlock] = []

        total_elements = len(doc.paragraphs) + sum(len(t.rows) for t in doc.tables)
        if total_elements > self.MAX_ELEMENTS:
            raise ValueError(
                f"DOCX exceeds maximum element limit: {total_elements} elements found (limit: {self.MAX_ELEMENTS})."
            )

        current_tier = "AUTHORITATIVE"
        current_reason = "docx_body"
        current_section = None
        block_idx = 0
        pseudo_page = 1
        elements_on_page = 0
        MAX_PER_PAGE = 35

        for element in doc.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

            if tag == "p":
                p = Paragraph(element, doc)
                text = p.text.strip()
                if not text:
                    continue

                # Section tracking & tier classification
                if re.search(r'\b(?:General\s+Terms|GTC|General\s+Conditions)\b', text, re.I):
                    current_tier = "CONTEXT_ONLY"
                    current_reason = "gtc_boilerplate"
                    current_section = text
                elif re.search(r'\b(?:Technical\s+Spec|Schedule\s+of\s+Requirements|Bill\s+of\s+Quantities|BoQ)\b', text, re.I):
                    current_tier = "AUTHORITATIVE"
                    current_reason = "technical_spec"
                    current_section = text
                elif re.search(r'\b(?:GeMARPTS|Searched\s+Result)\b', text, re.I):
                    current_tier = "CONTEXT_ONLY"
                    current_reason = "gemarpts_search_result"
                    current_section = text

                blocks.append(PageBlock(
                    page_no=pseudo_page,
                    block_idx=block_idx,
                    bbox=(0.0, float(block_idx * 20), 500.0, float(block_idx * 20 + 15)),
                    text=text,
                    raw_text=text,
                    tier=current_tier,
                    tier_reason=current_reason,
                    devanagari_ratio=0.0,
                    section_label=current_section
                ))
                block_idx += 1
                elements_on_page += 1

            elif tag == "tbl":
                table = Table(element, doc)
                for row_idx, row in enumerate(table.rows):
                    cell_texts = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    # Deduplicate merged cells that repeat identical text consecutively
                    deduped_cells = []
                    for c in cell_texts:
                        if not deduped_cells or c != deduped_cells[-1]:
                            deduped_cells.append(c)
                    row_line = " | ".join(deduped_cells).strip()
                    if not row_line or row_line == "|":
                        continue

                    blocks.append(PageBlock(
                        page_no=pseudo_page,
                        block_idx=block_idx,
                        bbox=(0.0, float(block_idx * 20), 500.0, float(block_idx * 20 + 15)),
                        text=row_line,
                        raw_text=row_line,
                        tier="AUTHORITATIVE",
                        tier_reason="table_row_authoritative",
                        devanagari_ratio=0.0,
                        section_label=current_section
                    ))
                    block_idx += 1
                    elements_on_page += 1

            if elements_on_page >= MAX_PER_PAGE:
                pseudo_page += 1
                elements_on_page = 0

        return blocks
