"""
build_docx.py
=============
Converts StandardSphere_Technical_Specification.md to a highly polished,
executive-grade Microsoft Word document (.docx).
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding (in twips) of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                       f'<w:top w:w="{top}" w:type="dxa"/>'
                       f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                       f'<w:left w:w="{left}" w:type="dxa"/>'
                       f'<w:right w:w="{right}" w:type="dxa"/>'
                       f'</w:tcMar>')
    tc_pr.append(tc_mar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Applies subtle borders to a table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tblBorders>')
    tbl_pr.append(borders)

def add_styled_paragraph(doc, text="", style='Normal', space_before=0, space_after=4, line_spacing=1.15):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    return p

def format_inline_runs(paragraph, text, base_font_size=10.5, is_italic=False):
    """
    Parses inline markdown: **bold**, *italic*, `code`, and $math$
    """
    # Regex to split on tokens
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$|__.*?__|_.*?_)')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = 'Calibri'
        run.font.size = Pt(base_font_size)
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
        
        if part.startswith('**') and part.endswith('**'):
            run.text = part[2:-2]
            run.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
        elif part.startswith('__') and part.endswith('__'):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run.text = part[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = RGBColor(194, 24, 91) # Crimson/Pink for code
            # Shading for inline code
            r_pr = run._r.get_or_add_rPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
            r_pr.append(shd)
        elif part.startswith('$') and part.endswith('$'):
            # Math
            run.text = part[1:-1]
            run.italic = True
            run.font.name = 'Cambria Math'
            run.font.color.rgb = RGBColor(30, 58, 138) # Dark Blue
        elif (part.startswith('*') and part.endswith('*')) or (part.startswith('_') and part.endswith('_')):
            run.text = part[1:-1]
            run.italic = True
        else:
            run.text = part
            if is_italic:
                run.italic = True

def add_code_block(doc, code_text):
    """Creates a boxed code/preformatted block with light background and border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only (accent)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="none"/>'
                        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="2563EB"/>'
                        f'<w:bottom w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tcBorders>')
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    for line in code_text.strip().split('\n'):
        run = p.add_run(line + '\n')
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(15, 23, 42)

def add_math_callout(doc, math_formula):
    """Renders a display math formula in a centered callout block."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    
    run = p.add_run(math_formula.strip())
    run.font.name = 'Cambria Math'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(30, 58, 138)
    run.italic = True

def add_callout(doc, text, title=""):
    """Styled callout note with a blue accent line."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="none"/>'
                        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="3B82F6"/>'
                        f'<w:bottom w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tcBorders>')
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    if title:
        r_title = p.add_run(title + "\n")
        r_title.bold = True
        r_title.font.name = 'Calibri'
        r_title.font.size = Pt(10.5)
        r_title.font.color.rgb = RGBColor(30, 64, 175)
    
    format_inline_runs(p, text, base_font_size=10)

def convert_markdown_to_docx(md_path, docx_path):
    print(f"Reading markdown from {md_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = docx.Document()
    
    # ── Margins ─────────────────────────────────────────────────────────────
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("StandardSphere: Technical Design Specification v2.0 | SIH-26108")
        hrun.font.name = 'Calibri'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Confidential — National Public Procurement Technical Blueprint")
        frun.font.name = 'Calibri'
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    lines = md_text.split('\n')
    i = 0
    total_lines = len(lines)
    
    in_code_block = False
    code_lines = []
    
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        
        # Parse table lines
        header_raw = table_lines[0]
        rows_raw = table_lines[2:] if len(table_lines) > 2 and '---' in table_lines[1] else table_lines[1:]
        
        headers = [c.strip() for c in header_raw.split('|')[1:-1]]
        num_cols = len(headers)
        if num_cols == 0:
            table_lines = []
            return
        
        tbl = doc.add_table(rows=len(rows_raw) + 1, cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True
        set_table_borders(tbl, color="CBD5E1", sz="4")
        
        # Style Header Row
        hdr_row = tbl.rows[0]
        # Repeat header on new page
        trPr = hdr_row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        
        for col_idx, h_text in enumerate(headers):
            cell = hdr_row.cells[col_idx]
            set_cell_background(cell, "1E3A8A") # Navy
            set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
            
        # Style Data Rows
        for r_idx, r_line in enumerate(rows_raw):
            row_cells = [c.strip() for c in r_line.split('|')[1:-1]]
            tbl_row = tbl.rows[r_idx + 1]
            bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
            
            for col_idx in range(num_cols):
                cell = tbl_row.cells[col_idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                
                val = row_cells[col_idx] if col_idx < len(row_cells) else ""
                format_inline_runs(p, val, base_font_size=9)
        
        # Spacer after table
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)
        table_lines = []

    while i < total_lines:
        line = lines[i]
        trimmed = line.strip()
        
        # ── Code Blocks (Fenced) ─────────────────────────────────────────────
        if trimmed.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                add_code_block(doc, '\n'.join(code_lines))
                code_lines = []
                i += 1
                continue
                
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # ── Tables (Markdown) ────────────────────────────────────────────────
        if trimmed.startswith('|') and trimmed.endswith('|'):
            table_lines.append(trimmed)
            i += 1
            continue
        else:
            if table_lines:
                flush_table()

        # ── Display Math ($$ ... $$) ─────────────────────────────────────────
        if trimmed.startswith('$$') and trimmed.endswith('$$') and len(trimmed) > 2:
            add_math_callout(doc, trimmed[2:-2])
            i += 1
            continue
        elif trimmed == '$$':
            math_buf = []
            i += 1
            while i < total_lines and lines[i].strip() != '$$':
                math_buf.append(lines[i].strip())
                i += 1
            add_math_callout(doc, ' '.join(math_buf))
            i += 1
            continue

        # ── Blank Line ───────────────────────────────────────────────────────
        if not trimmed:
            i += 1
            continue

        # ── Horizontal Rule ──────────────────────────────────────────────────
        if trimmed in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E2E8F0"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # ── Headings ─────────────────────────────────────────────────────────
        if trimmed.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(24)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(trimmed[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = RGBColor(15, 44, 89) # Deep Navy
            i += 1
            continue

        if trimmed.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(trimmed[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Navy Blue
            # Bottom border for H2
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="2" w:color="DBEAFE"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        if trimmed.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(trimmed[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 64, 175) # Blue 800
            i += 1
            continue

        if trimmed.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(trimmed[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
            i += 1
            continue

        # ── Unordered List ───────────────────────────────────────────────────
        if trimmed.startswith('* ') or trimmed.startswith('- '):
            indent_level = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            format_inline_runs(p, trimmed[2:])
            i += 1
            continue

        # ── Ordered List ─────────────────────────────────────────────────────
        m_num = re.match(r'^(\d+)\.\s+(.*)$', trimmed)
        if m_num:
            num = m_num.group(1)
            content = m_num.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.3)
            format_inline_runs(p, content)
            i += 1
            continue

        # ── Blockquote (>) ───────────────────────────────────────────────────
        if trimmed.startswith('> '):
            quote_text = trimmed[2:]
            add_callout(doc, quote_text)
            i += 1
            continue

        # ── Regular Paragraph ────────────────────────────────────────────────
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        format_inline_runs(p, trimmed)
        i += 1

    if table_lines:
        flush_table()

    print(f"Saving formatted Word document to {docx_path}...")
    doc.save(docx_path)
    print("Done! Document successfully generated.")

if __name__ == '__main__':
    md_file = r"C:\Users\devag\.gemini\antigravity\brain\9eb58232-caf4-49e0-b00d-7f324f33daf3\StandardSphere_Technical_Specification.md"
    docx_file = r"C:\Users\devag\.gemini\antigravity\brain\9eb58232-caf4-49e0-b00d-7f324f33daf3\StandardSphere_Technical_Specification.docx"
    convert_markdown_to_docx(md_file, docx_file)
