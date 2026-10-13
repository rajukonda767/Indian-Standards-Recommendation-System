# -*- coding: utf-8 -*-
"""
tests/fixtures/build_docx_xlsx_fixtures.py
=========================================
Builds structured test fixtures for DOCX and XLSX adapters:
1. tests/fixtures/sample_tender_spec.docx: Structured Word document with GTC (CONTEXT_ONLY)
   and Technical Specifications table (AUTHORITATIVE).
2. tests/fixtures/sample_boq_schedule.xlsx: Structured Excel BoQ workbook with headers,
   merged cells, quantities, units, and standard citations.
"""
import os
import docx
from docx.shared import Inches, Pt
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

FIXTURE_DIR = r"C:\Users\devag\.gemini\antigravity\scratch\standards\tests\fixtures"
os.makedirs(FIXTURE_DIR, exist_ok=True)

# ── 1. Construct Sample DOCX Tender Document ──
docx_path = os.path.join(FIXTURE_DIR, "sample_tender_spec.docx")
doc = docx.Document()

# Document Title
title_p = doc.add_heading("CENTRAL ELECTRICITY TENDER SPECIFICATION", level=0)

# Section 1: GTC Boilerplate (CONTEXT_ONLY)
doc.add_heading("Section 1: General Terms and Conditions (GTC)", level=1)
doc.add_paragraph("This procurement is governed by the General Conditions of Contract (GCC). All standard arbitration rules apply.")
doc.add_paragraph("GeMARPTS Search Result: reference IS 4989 foam concentrate generated during category search.")

# Section 2: Technical Specifications & BoQ Table (AUTHORITATIVE)
doc.add_heading("Section 2: Technical Specifications & Bill of Quantities", level=1)
doc.add_paragraph("All supplied electrical goods must strictly conform to the following technical schedule:")

table = doc.add_table(rows=1, cols=6)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
headers = ["Sl. No", "Material Code", "Item Description", "Technical Specification", "Quantity", "Unit"]
for i, h in enumerate(headers):
    hdr_cells[i].text = h

docx_items = [
    ("1", "CBL-101", "1.1 kV XLPE Power Cable 1C x 1.5 sq mm", "Manufactured and tested strictly as per IS 7098 (Part 1)", "5000", "MTR"),
    ("2", "CBL-102", "Heavy Duty Copper Control Cable 4C x 2.5", "Conforming to IS 1554 (Part 1):1988 with PVC outer sheath", "2500", "MTR"),
    ("3", "STL-201", "High Strength Deformed TMT Steel Bars Fe 500D", "Thermo-mechanically treated rebars per IS 1786:2008", "100", "MT"),
]

for item in docx_items:
    row_cells = table.add_row().cells
    for i, val in enumerate(item):
        row_cells[i].text = val

doc.save(docx_path)
print(f"Created constructed DOCX fixture: {docx_path} ({os.path.getsize(docx_path)} bytes)")


# ── 2. Construct Sample XLSX BoQ Schedule ──
xlsx_path = os.path.join(FIXTURE_DIR, "sample_boq_schedule.xlsx")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Schedule of Requirements"

# Header Banner
ws["A1"] = "PROJECT SCHEDULE OF REQUIREMENTS & BILL OF QUANTITIES"
ws.merge_cells("A1:F1")
ws["A1"].font = Font(bold=True, size=13)

# Column Headers
headers = ["Item No", "Material Code", "Item Description", "Technical Specification", "Estimated Quantity", "Unit"]
ws.append([])  # blank row 2
ws.append(headers)  # row 3

xlsx_items = [
    (1, "MAT-CBL-01", "1.1kV XLPE Armoured Power Cable 3.5C x 50 sq mm", "Specification: IS 7098 (Part 1) with extruded PVC outer sheath", 3000, "MTR"),
    (2, "MAT-CBL-02", "Single Core PVC Insulated Copper Building Wire 1.5 sq mm", "Conforming strictly to IS 694:2010 with ISI mark", 8000, "MTR"),
    (3, "MAT-PIP-01", "Mild Steel Heavy Grade Water Pipes 50mm NB", "Seamless steel tubes conforming to IS 1239 (Part 1):2004", 1500, "MTR"),
    (4, "MAT-CMT-01", "Ordinary Portland Cement 53 Grade", "High grade OPC conforming to IS 269:2015 in standard 50kg bags", 1000, "BAG"),
]

for row_data in xlsx_items:
    ws.append(list(row_data))

wb.save(xlsx_path)
print(f"Created constructed XLSX fixture: {xlsx_path} ({os.path.getsize(xlsx_path)} bytes)")
